"""生成两份 PRD（变更记录 / 字典管理），带截图，供导入飞书云文档。"""
from pathlib import Path
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn

BRANCH = Path('/Users/yuejun/work/project/ontology-studio-manual-branch')
ASSETS = BRANCH / 'docs' / 'prd-assets'
OUT = BRANCH / 'docs'
FONT = 'PingFang SC'
CN = 'PingFang SC'


def setup(doc):
    normal = doc.styles['Normal']
    normal.font.name = FONT
    normal.font.size = Pt(10.5)
    rpr = normal.element.get_or_add_rPr()
    rpr.get_or_add_rFonts().set(qn('w:eastAsia'), CN)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25
    for name in ('Heading 1', 'Heading 2', 'Heading 3', 'Title'):
        try:
            st = doc.styles[name]
        except KeyError:
            continue
        st.font.name = FONT
        st.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), CN)
        st.font.color.rgb = RGBColor(0x1F, 0x3A, 0x53)
    for section in doc.sections:
        section.top_margin = section.bottom_margin = Inches(0.7)
        section.left_margin = section.right_margin = Inches(0.8)


def p(doc, text, bold=False, size=None, italic=False, align=None):
    para = doc.add_paragraph()
    run = para.add_run(text)
    run.bold = bold
    run.italic = italic
    if size:
        run.font.size = Pt(size)
    if align:
        para.alignment = align
    return para


def bullets(doc, items):
    for item in items:
        doc.add_paragraph(item, style='List Bullet')


def numbers(doc, items):
    for item in items:
        doc.add_paragraph(item, style='List Number')


def table(doc, header, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(header))
    t.style = 'Light Grid Accent 1'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, text in enumerate(header):
        cell = t.rows[0].cells[i]
        cell.text = ''
        run = cell.paragraphs[0].add_run(text)
        run.bold = True
        run.font.size = Pt(9.5)
    for row in rows:
        cells = t.add_row().cells
        for i, text in enumerate(row):
            cells[i].text = ''
            run = cells[i].paragraphs[0].add_run(str(text))
            run.font.size = Pt(9.5)
            run.font.name = FONT
            run.element.rPr.rFonts.set(qn('w:eastAsia'), CN)
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Inches(w)
    doc.add_paragraph()
    return t


def figure(doc, filename, caption, width_in):
    path = ASSETS / filename
    if not path.exists():
        p(doc, f'（缺少截图：{filename}）', italic=True)
        return
    doc.add_picture(str(path), width=Inches(width_in))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = cap.add_run(caption)
    run.italic = True
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(0x6B, 0x7C, 0x8C)


def doc_info(doc, rows):
    table(doc, ['项目', '内容'], rows, widths=[1.4, 4.9])


# ===========================================================================
# PRD 1：变更记录
# ===========================================================================
d = Document()
setup(d)
d.add_heading('变更记录 · 产品需求文档', level=0)
p(d, '本体中心 → 变更与发布 → 变更记录', italic=True)
doc_info(d, [
    ['功能名称', '变更记录'],
    ['所属模块', '本体中心 → 变更与发布'],
    ['页面路由', '#/ontology/governance/activity'],
    ['适用对象', '本体建模人员、数据治理人员、评审人'],
    ['文档状态', '评审中'],
    ['更新日期', '2026-10-12'],
])

d.add_heading('1. 功能概述', level=1)
p(d, '变更记录回答一个问题：这个本体上，谁在什么时候，把哪个资源，从什么改成了什么。')
p(d, '它只展示当前本体的操作流水，不与其他本体的记录混排；每条记录都定位到一个具体资源，'
     '并支持在「记录」与「资源」之间双向跳转——看到变更就能直接跳过去修，修完还能回到原来那条记录。')

d.add_heading('2. 目标与使用场景', level=1)
bullets(d, [
    '追溯：发布前核对「这次到底改了什么」，而不是只看一句操作摘要。',
    '定位：从一条变更直接跳到出问题的对象、关系或映射去修改。',
    '审计：按人、时间、资源类型回溯建模过程，并可选择是否包含系统自动操作。',
    '交接：新成员通过变更流理解模型是如何一步步演进出来的。',
])

d.add_heading('3. 页面结构', level=1)
p(d, '页面为「左侧列表 + 右侧详情」两栏：列表用于快速扫读与筛选，详情用于看清一条记录的全部证据。')
figure(d, 'A1-变更记录-整页.png', '图 1　变更记录整页：列表、筛选工具条与右侧详情', 6.4)

d.add_heading('4. 列表区', level=1)
d.add_heading('4.1 列说明', level=2)
table(d, ['列', '内容'], [
    ['时间', '操作时间，精确到秒'],
    ['资源', '资源名称 + 资源类型 + 资源 ID（ID 可复制用于排查）'],
    ['操作', '新增 / 修改 / 删除；发布等其它操作按原始操作名展示'],
    ['变更内容', '变更摘要；有字段级差异时直接显示「字段：旧 → 新」'],
    ['操作人', '显示为可读名称，例如「手工建模管理员」'],
], widths=[1.1, 5.2])

d.add_heading('4.2 筛选与搜索', level=2)
figure(d, 'A4-变更记录-筛选工具条.png', '图 2　筛选工具条', 6.4)
bullets(d, [
    '关键词：可搜索资源名称、资源 ID、变更内容、操作人。',
    '资源类型：24 类，包括对象类型、关系类型、动作类型、接口、规则、数据源、治理资产、数据映射等。',
    '操作：全部 / 新增 / 修改 / 删除。',
    '包含系统操作：默认关闭；勾选后显示系统自动产生的记录，便于把人工操作与系统操作分开看。',
    '分页：默认每页 20 条，可切换；右上角显示当前筛选下的总条数。',
])

d.add_heading('5. 详情区', level=1)
figure(d, 'A2-变更记录-详情面板.png', '图 3　变更详情：资源定位、操作信息、变更内容与字段变化', 2.3)
d.add_heading('5.1 头部与资源定位', level=2)
bullets(d, [
    '头部：操作类型徽标（新增/修改/删除）+ 资源名称 + 资源类型 + 「打开该资源」按钮。',
    '资源定位：记录编号、资源类型、资源 ID、所属本体。',
    '操作信息：操作人、操作时间。',
])
d.add_heading('5.2 变更内容', level=2)
p(d, '逐行展示本次变更的摘要。两侧相同的「空变化」（例如旧值与新值都是空）不会展示，避免出现“无 → 无”这类没有信息量的行。')
d.add_heading('5.3 字段变化 · 操作时快照', level=2)
figure(d, 'A3-变更记录-字段变化.png', '图 4　字段级差异：修改前 / 修改后', 3.0)
bullets(d, [
    '以「字段 + 修改前 + 修改后」的形式呈现。',
    '属性类变化会细化到属性级，例如「属性 · 客户分类（fl）」。',
    '删除类记录额外展示「删除前资源摘要」，保留删除时刻的资源状态。',
])

d.add_heading('6. 资源往返跳转', level=1)
table(d, ['方向', '行为'], [
    ['记录 → 资源', '点「打开该资源」跳转到该资源的编辑或查看页，并自动选中该资源'],
    ['资源 → 记录', '各资源页详情里的「变更记录」入口，进入该资源的专属记录视图'],
    ['返回', '从资源页回到变更记录时，恢复原来的记录选中与页码'],
    ['资源已删除', '提示「该资源已不存在或不属于当前本体」，可关闭提示；历史记录仍然保留'],
    ['无独立页面', '旧版对象类型组、共享属性、智能构建提案、建模证据、属性等，明确说明无法跳转，不给无效按钮'],
], widths=[1.2, 5.1])

d.add_heading('7. 规则与口径', level=1)
table(d, ['规则', '说明'], [
    ['记录范围', '仅当前本体；不做跨本体混排'],
    ['字段级差异的产生', '修改类操作保存操作前后的资源快照并逐字段比较；新增与删除保存单边快照'],
    ['比较范围', '只比较语义字段；ID、创建/更新时间、密钥等内部字段不参与'],
    ['空值口径', 'None 与空字符串都表示「未设置」，不产生「无 → 无」这类空变化'],
    ['属性变化粒度', '细化到 properties.<字段>，例如「属性 · 交付日期（deliveryDate）」'],
    ['覆盖资源', '本体、业务域、主题域、对象类型、接口、关系类型、动作类型、业务规则、数据源、治理资产、数据映射、业务场景、系统能力、动作执行映射、权限策略、规则与函数'],
    ['历史记录', '功能上线前写入的记录可能只有文本摘要，没有结构化字段差异；系统不推断未记录的旧值'],
], widths=[1.4, 4.9])

d.add_heading('8. 边界与已知限制', level=1)
bullets(d, [
    '实例数据的增删改目前只记录操作摘要，尚未提供字段级差异。',
    '发布、智能构建提案合并等操作按原始操作名展示，不归入新增/修改/删除三类。',
    '变更记录是只读审计视图，不提供「回滚到某条记录」。',
    '不展示实例数据内容与生产接口运行结果。',
])

d.add_heading('9. 验收要点', level=1)
bullets(d, [
    '列表只显示当前本体的记录，切换本体后记录随之切换。',
    '修改类记录能看到「字段：旧 → 新」，属性变化细化到属性级。',
    '空值变化（None 与空字符串互换）不出现「无 → 无」。',
    '「打开该资源」能跳到对应资源并选中，再跳回来仍选中同一条记录。',
    '删除类资源跳转时给出明确提示，而不是空白页。',
    '关键词、资源类型、操作、包含系统操作四项筛选可组合，切换筛选时不丢作用域。',
    '分页与总条数一致。',
])

d.save(OUT / 'PRD-变更记录.docx')
print('已生成 PRD-变更记录.docx')

# ===========================================================================
# PRD 2：字典管理
# ===========================================================================
d = Document()
setup(d)
d.add_heading('字典管理 · 产品需求文档', level=0)
p(d, '本体中心 → 本体模型 → 字典管理', italic=True)
doc_info(d, [
    ['功能名称', '字典管理（全局共享语义选项）'],
    ['所属模块', '本体中心 → 本体模型'],
    ['页面路由', '#/ontology/model/dictionaries'],
    ['适用对象', '本体建模人员、数据治理人员'],
    ['文档状态', '评审中'],
    ['更新日期', '2026-10-12'],
])

d.add_heading('1. 功能概述', level=1)
p(d, '字典管理维护全局共享的建模选项，让同一个概念在所有本体里写法一致——避免一个本体写「直接映射」、'
     '另一个本体写「direct」。')
p(d, '业务域与主题域属于本体目录实体，不在字典里管理；它们有独立的目录页面。')

d.add_heading('2. 字典类别', level=1)
p(d, '系统统一管理 7 类全局字典：')
table(d, ['类别', '代码', '典型字典项'], [
    ['数据类型', 'data_type', '字符串 string、整数 integer、小数 decimal、布尔 boolean、日期时间 datetime、枚举 enum'],
    ['系统类型', 'system_type', '业务系统 business、实时系统 realtime、资产系统 asset、财务系统 finance、监测系统 monitor、调度系统 dispatch'],
    ['数据类别', 'data_category', '主数据、交易数据、时序数据、分析数据 等'],
    ['接口方式', 'interface_type', '数据库、接口、REST API、文件交换、消息队列 等'],
    ['映射方式', 'mapping_way', '直接映射、表达式转换、枚举映射 等'],
    ['行动类型', 'action_type', '业务行动、通知行动、控制行动 等'],
    ['转换规则', 'xform_rule', '单位标准化、状态值归一 等'],
], widths=[1.1, 1.3, 3.9])

d.add_heading('3. 页面结构', level=1)
p(d, '页面为「左侧类别导航 + 右侧（列表 + 详情）」结构：左侧切换字典类别，右侧维护该类别下的字典项。')
figure(d, 'B1-字典管理-整页.png', '图 1　字典管理整页：左侧 7 类字典，右侧字典项列表与详情', 6.4)

d.add_heading('4. 类别导航与字典项列表', level=1)
figure(d, 'B5-字典管理-行动类型.png', '图 2　切换到「行动类型」类别', 6.4)
bullets(d, [
    '左侧 7 个类别入口，每项实时显示该类别的字典项数量。',
    '列表列：名称 / 代码、状态（有效 / 禁用）、排序值。',
    '工具条：按名称或代码搜索、新增字典项、刷新。',
])

d.add_heading('5. 字典项详情与编辑', level=1)
figure(d, 'B3-字典管理-字典项详情.png', '图 3　字典项详情', 1.9)
figure(d, 'B4-字典管理-新增表单.png', '图 4　新增 / 编辑字典项', 1.9)
p(d, '通用字段：代码、名称、描述、排序、状态。')
p(d, '类别特有字段：数据类型可维护「默认值」「系统代码」；系统类型可维护「联络人」。')
bullets(d, [
    '代码：创建后不可修改，它是被模型引用的稳定标识。',
    '状态：有效 / 禁用。禁用项不再出现在新建下拉中，但历史值仍可查看。',
    '编辑：详情右上角铅笔图标；删除：垃圾桶图标，删除前二次确认。',
])

d.add_heading('6. 表单内联管理', level=1)
p(d, '在数据源、数据映射、对象属性、行动等表单中，每个字典下拉旁都有「管理此字典」，点开即为就地管理面板，'
     '可以不离开当前表单完成增删改；新建的字典项会自动写回当前表单，省去"保存表单 → 去字典页新增 → 回来重填"的往返。')
figure(d, 'C1-表单内联字典管理.png', '图 5　登记数据源表单里的内联字典管理（系统类型）', 5.6)

d.add_heading('7. 规则与口径', level=1)
table(d, ['规则', '说明'], [
    ['代码不可改', '代码是稳定引用标识，创建后不能修改'],
    ['在用项不可禁用或删除', '服务端校验并返回 409，提示“字典项已被模型引用，不能删除 / 禁用；请先处理引用”'],
    ['名称被引用时不可改名', '历史数据中存在按名称保存的引用，直接改名会断链；服务端返回 409 并提示先完成引用迁移'],
    ['禁用项不出现在新建下拉', '已有历史值仍可在详情与下拉中看到，标注为「已禁用，历史值」'],
    ['行为类新增默认禁用', '数据类型、接口方式、映射方式、行动类型、转换规则新增时先保存为禁用，接入处理能力后再启用'],
    ['字典外历史值保留', '已保存但不在字典中的取值显示为「（已有值，待规范）」，不静默丢弃'],
    ['排序', '数值排序，用于控制下拉中字典项的先后顺序'],
], widths=[1.5, 4.8])

d.add_heading('8. 与建模的关系', level=1)
bullets(d, [
    '数据类型：影响对象属性与规则参数的类型定义。',
    '系统类型、数据类别、接口方式、转换规则：影响数据源登记。',
    '映射方式：影响语义映射的转换配置。',
    '行动类型：影响动作类型的定义。',
    '字典项的增删改会立即刷新所有使用该字典的下拉选项。',
])

d.add_heading('9. 边界与已知限制', level=1)
bullets(d, [
    '同一字段存在「按代码保存」与「按名称保存」两种历史写法，尚未统一迁移；这也是改名守卫采用保守策略的原因。',
    '暂未对字典外取值做强制校验，也没有「字典外取值清单」用于批量归一。',
    '字典是全局资源，不属于任何本体，因此字典项的编辑目前不写入本体的变更记录。',
    '内联面板中对当前选中项改名或删除时，外层表单不会自动回写取值，需要重新选择。',
])

d.add_heading('10. 验收要点', level=1)
bullets(d, [
    '7 类字典均可增删改查，代码创建后不可修改。',
    '在用字典项的禁用与删除被服务端拒绝，并给出可读原因。',
    '按名称被引用的字典项改名被拒绝，并提示先处理引用。',
    '禁用项不出现在新建下拉，历史值仍可见。',
    '表单内联管理可新增字典项并自动写回当前表单。',
    '左侧类别导航的项数随增删实时变化。',
])

d.save(OUT / 'PRD-字典管理.docx')
print('已生成 PRD-字典管理.docx')
