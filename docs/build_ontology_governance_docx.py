"""把《本体治理方案》Markdown 转成排版后的 DOCX。

用法：python3 docs/build_ontology_governance_docx.py
"""
from __future__ import annotations

import re
import unicodedata
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

BASE = Path(__file__).resolve().parent
SRC = BASE / "本体治理方案-通用方法论与电视台一期实践.md"
OUT = BASE / "本体治理方案-通用方法论与电视台一期实践.docx"

INK = RGBColor(0x1F, 0x3A, 0x53)
INK_LIGHT = RGBColor(0x44, 0x58, 0x6B)
MUTED = RGBColor(0x6B, 0x7A, 0x88)
ACCENT = RGBColor(0x1F, 0x6F, 0x8B)
FONT_CN = "PingFang SC"
FONT_MONO = "Menlo"

TEXT_WIDTH_CM = 17.0


# ---------- 基础工具 ----------

def set_run_font(run, *, size=None, bold=None, color=None, font=FONT_CN, italic=None):
    run.font.name = font
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:ascii"), font)
    rfonts.set(qn("w:hAnsi"), font)
    rfonts.set(qn("w:eastAsia"), FONT_CN if font == FONT_CN else FONT_CN)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = color


INLINE = re.compile(r"(\*\*.+?\*\*|`[^`]+`)")


def add_inline(paragraph, text, *, size=10.5, color=None, bold=False):
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            run = paragraph.add_run(part[2:-2])
            set_run_font(run, size=size, bold=True, color=color)
        elif part.startswith("`") and part.endswith("`") and len(part) > 2:
            run = paragraph.add_run(part[1:-1])
            set_run_font(run, size=size - 1, font=FONT_MONO, color=ACCENT)
        else:
            run = paragraph.add_run(part)
            set_run_font(run, size=size, bold=bold, color=color)


TCPR_ORDER = [
    "w:cnfStyle", "w:tcW", "w:gridSpan", "w:hMerge", "w:vMerge", "w:tcBorders",
    "w:shd", "w:noWrap", "w:tcMar", "w:textDirection", "w:tcFitText",
    "w:vAlign", "w:hideMark", "w:headers", "w:cellIns", "w:cellDel",
    "w:cellMerge", "w:tcPrChange",
]

PPR_ORDER = [
    "w:pStyle", "w:keepNext", "w:keepLines", "w:pageBreakBefore", "w:framePr",
    "w:widowControl", "w:numPr", "w:suppressLineNumbers", "w:pBdr", "w:shd",
    "w:tabs", "w:suppressAutoHyphens", "w:kinsoku", "w:wordWrap",
    "w:overflowPunct", "w:topLinePunct", "w:autoSpaceDE", "w:autoSpaceDN",
    "w:bidi", "w:adjustRightInd", "w:snapToGrid", "w:spacing", "w:ind",
    "w:contextualSpacing", "w:mirrorIndents", "w:suppressOverlap", "w:jc",
    "w:textDirection", "w:textAlignment", "w:textboxTightWrap",
    "w:outlineLvl", "w:divId", "w:cnfStyle", "w:rPr", "w:sectPr",
    "w:pPrChange",
]

TBLPR_ORDER = [
    "w:tblStyle", "w:tblpPr", "w:tblOverlap", "w:bidiVisual",
    "w:tblStyleRowBandSize", "w:tblStyleColBandSize", "w:tblW", "w:jc",
    "w:tblCellSpacing", "w:tblInd", "w:tblBorders", "w:shd", "w:tblLayout",
    "w:tblCellMar", "w:tblLook", "w:tblCaption", "w:tblDescription",
]


def insert_ordered(parent, child, order):
    tag = child.tag
    index = order.index(tag) if tag in order else len(order)
    for existing in list(parent):
        if existing.tag in order and order.index(existing.tag) > index:
            existing.addprevious(child)
            return
    parent.append(child)


def shade(element, fill, order=PPR_ORDER):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    insert_ordered(element, shd, order)


def set_grid(table, widths):
    for index, width in enumerate(widths):
        table.columns[index].width = width
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            cell.width = widths[index]


def display_width(text: str) -> int:
    width = 0
    for ch in text:
        width += 2 if unicodedata.east_asian_width(ch) in ("W", "F") else 1
    return width


def para_border(paragraph, *, left=None, bottom=None):
    ppr = paragraph._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    if left:
        el = OxmlElement("w:left")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(left[1]))
        el.set(qn("w:space"), "6")
        el.set(qn("w:color"), left[0])
        pbdr.append(el)
    if bottom:
        el = OxmlElement("w:bottom")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(bottom[1]))
        el.set(qn("w:space"), "4")
        el.set(qn("w:color"), bottom[0])
        pbdr.append(el)
    insert_ordered(ppr, pbdr, PPR_ORDER)


# ---------- 文档骨架 ----------

doc = Document()
doc.core_properties.title = "本体治理方案：从数据标准到可执行的语义契约"
doc.core_properties.subject = "通用方法论与北京电视台一期落地实践"
doc.core_properties.author = "企业本体 Studio"
doc.core_properties.comments = "本体驱动的数据治理实施方案 v1.0"

section = doc.sections[0]
section.page_width = Cm(21.0)
section.page_height = Cm(29.7)
section.left_margin = Cm(2.2)
section.right_margin = Cm(1.8)
section.top_margin = Cm(2.2)
section.bottom_margin = Cm(2.0)
section.header_distance = Cm(1.2)
section.footer_distance = Cm(1.0)
section.different_first_page_header_footer = True

normal = doc.styles["Normal"]
normal.font.name = FONT_CN
normal.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), FONT_CN)
normal.font.size = Pt(10.5)
normal.font.color.rgb = RGBColor(0x22, 0x2A, 0x33)
normal.paragraph_format.line_spacing = 1.32
normal.paragraph_format.space_after = Pt(6)

heading_specs = {
    "Heading 1": (16, 18, 8),
    "Heading 2": (12.5, 14, 5),
    "Heading 3": (11, 10, 4),
}
for name, (size, before, after) in heading_specs.items():
    st = doc.styles[name]
    st.font.name = FONT_CN
    st.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), FONT_CN)
    st.font.size = Pt(size)
    st.font.bold = True
    st.font.color.rgb = INK
    st.paragraph_format.space_before = Pt(before)
    st.paragraph_format.space_after = Pt(after)
    st.paragraph_format.keep_with_next = True
    st.paragraph_format.line_spacing = 1.15

bullet = doc.styles["List Bullet"]
bullet.font.name = FONT_CN
bullet.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), FONT_CN)
bullet.font.size = Pt(10.5)
bullet.paragraph_format.space_after = Pt(3)
bullet.paragraph_format.line_spacing = 1.28


# ---------- 页眉页脚 ----------

def add_page_field(paragraph):
    run = paragraph.add_run("第 ")
    set_run_font(run, size=8.5, color=MUTED)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run2 = paragraph.add_run()
    set_run_font(run2, size=8.5, color=MUTED)
    run2._r.append(begin)
    run2._r.append(instr)
    run2._r.append(end)
    run3 = paragraph.add_run(" 页")
    set_run_font(run3, size=8.5, color=MUTED)


header = section.header.paragraphs[0]
header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
h_run = header.add_run("本体治理方案  ｜  通用方法论与电视台一期实践")
set_run_font(h_run, size=8, color=MUTED)
para_border(header, bottom=("D5DDE5", 4))

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
f_run = footer.add_run("企业本体 Studio  ｜  v1.0  ｜  ")
set_run_font(f_run, size=8.5, color=MUTED)
add_page_field(footer)


# ---------- 封面 ----------

for _ in range(3):
    doc.add_paragraph()

kicker = doc.add_paragraph()
kicker_run = kicker.add_run("数 据 治 理 实 施 方 案")
set_run_font(kicker_run, size=11, bold=True, color=ACCENT)

title = doc.add_paragraph()
title.paragraph_format.space_before = Pt(10)
title.paragraph_format.space_after = Pt(4)
title_run = title.add_run("本体治理方案")
set_run_font(title_run, size=34, bold=True, color=INK)

subtitle = doc.add_paragraph()
sub_run = subtitle.add_run("从数据标准到可执行的语义契约")
set_run_font(sub_run, size=17, color=INK_LIGHT)

tagline = doc.add_paragraph()
tagline.paragraph_format.space_before = Pt(2)
tagline_run = tagline.add_run("通用方法论  ·  北京电视台一期落地实践")
set_run_font(tagline_run, size=12, color=MUTED)

rule = doc.add_paragraph()
rule.paragraph_format.space_before = Pt(16)
rule.paragraph_format.space_after = Pt(18)
para_border(rule, bottom=("1F3A53", 12))

quote = doc.add_paragraph()
quote.paragraph_format.left_indent = Cm(0.2)
quote.paragraph_format.right_indent = Cm(1.0)
quote_run = quote.add_run(
    "标准只有被系统执行、被数据验证、被业务和智能体消费，才算真正落地。"
)
set_run_font(quote_run, size=12.5, color=INK, italic=False)

for _ in range(5):
    doc.add_paragraph()

meta = doc.add_table(rows=4, cols=2)
meta.alignment = WD_TABLE_ALIGNMENT.LEFT
meta.autofit = False
meta_rows = [
    ("文档版本", "v1.0"),
    ("编制日期", "2026 年 10 月"),
    ("适用对象", "数据治理与本体建设项目 · 售前方案 / 启动交底"),
    ("使用说明", "客户侧系统名称与字段为待核实的行业通行假设，实施前须以调研结论为准"),
]
for idx, (k, v) in enumerate(meta_rows):
    cells = meta.rows[idx].cells
    cells[0].width = Cm(3.2)
    cells[1].width = Cm(13.0)
    p0 = cells[0].paragraphs[0]
    r0 = p0.add_run(k)
    set_run_font(r0, size=9, bold=True, color=MUTED)
    p0.paragraph_format.space_after = Pt(2)
    p1 = cells[1].paragraphs[0]
    r1 = p1.add_run(v)
    set_run_font(r1, size=9, color=INK_LIGHT)
    p1.paragraph_format.space_after = Pt(2)
    for cell in cells:
        tc_pr = cell._tc.get_or_add_tcPr()
        borders = OxmlElement("w:tcBorders")
        for edge in ("top", "left", "bottom", "right"):
            el = OxmlElement(f"w:{edge}")
            el.set(qn("w:val"), "nil")
            borders.append(el)
        tc_pr.append(borders)

set_grid(meta, [Cm(3.2), Cm(13.2)])

doc.add_page_break()


# ---------- 结构导览 ----------

lead = doc.add_paragraph()
lead_run = lead.add_run("内容结构")
set_run_font(lead_run, size=16, bold=True, color=INK)
lead.paragraph_format.space_after = Pt(10)
para_border(lead, bottom=("1F3A53", 8))

structure = [
    ("摘要", "一份判断、三个不等式、一张架构图、四件可验收资产与人机分工"),
    ("1  问题与判断", "为什么传统治理建了标准却落不了地：三个不等式与四个断点"),
    ("2  本体治理的定义", "三层架构、标准的四种身份、十类标准如何落在本体"),
    ("3  通用方法论", "五步闭环、四条流水线、组织角色、度量体系与六个失败模式"),
    ("4  治理智能体", "五种工作模式、工具集、输出契约、人工卡点与防幻觉规则"),
    ("5  工作流落地", "四层结构、技能设计、五条标准工作流、编排模式、反模式与灰度"),
    ("6  通用实施路线", "12 周计划、关键路径、团队投入与 90 天必须看到的成果"),
    ("7  电视台一期落地", "现状假设、目标边界、本体蓝图、三条主链、智能体用法、14 周计划与验收"),
    ("8  立即可做的五件事", "签约后两周内可启动的最小动作"),
    ("附录", "标准卡片模板、对象建模检查清单、任务卡模板、术语表"),
]

st = doc.add_table(rows=len(structure), cols=2)
st.autofit = False
st.alignment = WD_TABLE_ALIGNMENT.CENTER
for idx, (k, v) in enumerate(structure):
    cells = st.rows[idx].cells
    cells[0].width = Cm(4.2)
    cells[1].width = Cm(12.6)
    p0 = cells[0].paragraphs[0]
    r0 = p0.add_run(k)
    set_run_font(r0, size=10, bold=True, color=INK)
    p0.paragraph_format.space_before = Pt(4)
    p0.paragraph_format.space_after = Pt(4)
    p1 = cells[1].paragraphs[0]
    r1 = p1.add_run(v)
    set_run_font(r1, size=9.5, color=INK_LIGHT)
    p1.paragraph_format.space_before = Pt(4)
    p1.paragraph_format.space_after = Pt(4)
    tc_pr = cells[1]._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    el = OxmlElement("w:bottom")
    el.set(qn("w:val"), "single")
    el.set(qn("w:sz"), "4")
    el.set(qn("w:color"), "DDE4EA")
    borders.append(el)
    tc_pr.append(borders)

set_grid(st, [Cm(4.2), Cm(12.6)])

doc.add_page_break()


# ---------- 正文解析 ----------

raw = SRC.read_text(encoding="utf-8").splitlines()
start = next(i for i, line in enumerate(raw) if line.startswith("## 摘要"))
lines = raw[start:]


DIAGRAMS = [
    (("语义层（本体）", "物理层（数据源）", "运行层（消费）"),
     "arch-three-layers.png", Cm(15.4), "图 1　数据治理的三层架构：每一层承载标准的不同身份"),
    (("消费层", "源系统", "物理层"),
     "arch-four-layers.png", Cm(15.0), "图 2　四层治理架构：向下落标，向上服务，横向校验"),
    (("治理智能体（五种工作模式）",),
     "agent-architecture.png", Cm(15.0), "图 3　治理智能体的运行架构与发布前的四道关卡"),
    (("接入层", "编排层", "技能层"),
     "workflow-layers.png", Cm(15.2), "图 4　工作流智能体的四层结构与横切能力"),
]


def add_code_block(block_lines):
    body = "\n" + "\n".join(block_lines) + "\n"
    widest = max((display_width(l) for l in block_lines), default=0)
    size = 8.0 if widest <= 62 else 7.2 if widest <= 78 else 6.6
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.left_indent = Cm(0.35)
    paragraph.paragraph_format.right_indent = Cm(0.2)
    paragraph.paragraph_format.space_before = Pt(5)
    paragraph.paragraph_format.space_after = Pt(8)
    paragraph.paragraph_format.line_spacing = 1.12
    paragraph.paragraph_format.keep_together = True
    ppr = paragraph._p.get_or_add_pPr()
    shade(ppr, "F5F7FA")
    para_border(paragraph, left=("AFC2D2", 12))
    run = paragraph.add_run(body)
    set_run_font(run, size=size, font=FONT_MONO, color=RGBColor(0x2B, 0x38, 0x45))
    return paragraph


def split_row(line):
    cells = [c.strip() for c in line.strip().strip("|").split("|")]
    return cells


def is_separator(line):
    return bool(re.fullmatch(r"\|[\s:\-|]+\|", line.strip()))


def add_table(rows):
    header = split_row(rows[0])
    body = [split_row(r) for r in rows[2:]]
    ncols = len(header)
    table = doc.add_table(rows=1 + len(body), cols=ncols)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    weights = []
    for c in range(ncols):
        samples = [header[c]] + [r[c] if c < len(r) else "" for r in body]
        longest = max(display_width(s) for s in samples)
        weights.append(min(34, max(8, longest + 2)))
    total = sum(weights)
    widths_cm = [TEXT_WIDTH_CM * w / total for w in weights]
    min_first = 3.0 if ncols <= 3 else 2.6
    if header and header[0].strip() in ("序", "#", "序号"):
        min_first = 1.3
    if widths_cm[0] < min_first:
        deficit = min_first - widths_cm[0]
        rest = sum(widths_cm[1:]) or 1
        widths_cm = [min_first] + [
            value - deficit * value / rest for value in widths_cm[1:]
        ]
    widths = [Cm(value) for value in widths_cm]

    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:color"), "C9D3DB")
        borders.append(el)
    insert_ordered(tbl_pr, borders, TBLPR_ORDER)

    margins = OxmlElement("w:tblCellMar")
    for edge, val in (("top", 50), ("bottom", 50), ("left", 90), ("right", 90)):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:w"), str(val))
        el.set(qn("w:type"), "dxa")
        margins.append(el)
    insert_ordered(tbl_pr, margins, TBLPR_ORDER)

    set_grid(table, widths)

    font_size = 9.0 if ncols <= 4 else 8.4 if ncols == 5 else 7.8

    def fill(cell, text, *, bold=False, header_cell=False):
        cell.width = widths[cell_idx[0]]
        paragraph = cell.paragraphs[0]
        paragraph.paragraph_format.space_before = Pt(1)
        paragraph.paragraph_format.space_after = Pt(1)
        paragraph.paragraph_format.line_spacing = 1.18
        add_inline(paragraph, text, size=font_size, color=INK if header_cell else None, bold=bold)

    for c, text in enumerate(header):
        cell_idx = [c]
        fill(table.rows[0].cells[c], text, bold=True, header_cell=True)
        tc_pr = table.rows[0].cells[c]._tc.get_or_add_tcPr()
        shade(tc_pr, "E6EDF3", TCPR_ORDER)
    tr_pr = table.rows[0]._tr.get_or_add_trPr()
    tr_pr.append(OxmlElement("w:tblHeader"))

    for r, row in enumerate(body, start=1):
        for c in range(ncols):
            cell_idx = [c]
            text = row[c] if c < len(row) else ""
            fill(table.rows[r].cells[c], text)
        tr_pr = table.rows[r]._tr.get_or_add_trPr()
        tr_pr.append(OxmlElement("w:cantSplit"))

    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_before = Pt(0)
    spacer.paragraph_format.space_after = Pt(0)
    spacer.paragraph_format.line_spacing = 0.6
    return table


i = 0
while i < len(lines):
    line = lines[i]
    stripped = line.strip()

    if stripped == "" or stripped == "---":
        i += 1
        continue

    if stripped.startswith("```"):
        i += 1
        block = []
        while i < len(lines) and not lines[i].strip().startswith("```"):
            block.append(lines[i])
            i += 1
        i += 1
        block_text = "\n".join(block)
        matched = next(
            (item for item in DIAGRAMS if all(key in block_text for key in item[0])),
            None,
        )
        if matched:
            _, filename, width, caption = matched
            doc.add_picture(str(BASE / "diagram-assets" / filename), width=width)
            doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            doc.paragraphs[-1].paragraph_format.space_before = Pt(6)
            doc.paragraphs[-1].paragraph_format.space_after = Pt(2)
            doc.paragraphs[-1].paragraph_format.keep_with_next = True
            cap = doc.add_paragraph()
            cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            cap.paragraph_format.space_after = Pt(10)
            cap_run = cap.add_run(caption)
            set_run_font(cap_run, size=8.5, color=MUTED)
        else:
            add_code_block(block)
        continue

    if stripped.startswith("|"):
        rows = []
        while i < len(lines) and lines[i].strip().startswith("|"):
            rows.append(lines[i])
            i += 1
        if len(rows) >= 3 and is_separator(rows[1]):
            add_table(rows)
        continue

    if stripped.startswith("#### "):
        doc.add_heading(stripped[5:], level=3)
        i += 1
        continue
    if stripped.startswith("### "):
        doc.add_heading(stripped[4:], level=2)
        i += 1
        continue
    if stripped.startswith("## "):
        doc.add_heading(stripped[3:], level=1)
        i += 1
        continue
    if stripped.startswith("# "):
        i += 1
        continue

    if stripped.startswith(">"):
        text = stripped.lstrip(">").strip()
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.left_indent = Cm(0.4)
        paragraph.paragraph_format.right_indent = Cm(0.6)
        paragraph.paragraph_format.space_before = Pt(4)
        paragraph.paragraph_format.space_after = Pt(8)
        paragraph.paragraph_format.line_spacing = 1.3
        para_border(paragraph, left=("1F6F8B", 18))
        shade(paragraph._p.get_or_add_pPr(), "F3F6F8")
        add_inline(paragraph, text, size=10, color=INK_LIGHT)
        i += 1
        continue

    if stripped.startswith("- [ ] "):
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.left_indent = Cm(0.75)
        paragraph.paragraph_format.first_line_indent = Cm(-0.45)
        paragraph.paragraph_format.space_after = Pt(3)
        run = paragraph.add_run("☐  ")
        set_run_font(run, size=11, color=ACCENT)
        add_inline(paragraph, stripped[6:])
        i += 1
        continue

    if stripped.startswith("- "):
        paragraph = doc.add_paragraph(style="List Bullet")
        paragraph.paragraph_format.left_indent = Cm(0.75)
        paragraph.paragraph_format.first_line_indent = Cm(-0.35)
        add_inline(paragraph, stripped[2:])
        i += 1
        continue

    match = re.match(r"^(\d+)\.\s+(.*)$", stripped)
    if match:
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.left_indent = Cm(0.75)
        paragraph.paragraph_format.first_line_indent = Cm(-0.55)
        paragraph.paragraph_format.space_after = Pt(4)
        number_run = paragraph.add_run(f"{match.group(1)}.")
        set_run_font(number_run, size=10.5, bold=True, color=ACCENT)
        space_run = paragraph.add_run("  ")
        set_run_font(space_run, size=10.5)
        add_inline(paragraph, match.group(2))
        i += 1
        continue

    paragraph = doc.add_paragraph()
    add_inline(paragraph, stripped)
    i += 1

doc.save(OUT)
print(f"saved: {OUT}")
